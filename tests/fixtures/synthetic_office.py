"""
Generates fabricated Office/HTML/Markdown documents for tests (Phase 9,
master directive numbering) -- no real business/personal data, matching
every other fixture in this project (see docs/PHASE_0_AUDIT.md Risk
Register #1).

python-docx/python-pptx are dev-only (requirements-dev.txt): they build
these fixtures, they are not part of the runtime conversion path at all
-- that's LibreOffice, via documents_converter/converters/_libreoffice.py.
"""

from __future__ import annotations

from pathlib import Path

MARKER_TEXT = "Fabricated test document for Phase 9 conversion tests."


def build_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(MARKER_TEXT)
    doc.save(str(path))


def build_xlsx(path: str) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = MARKER_TEXT
    wb.save(str(path))


def build_pptx(path: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    textbox.text_frame.text = MARKER_TEXT
    prs.save(str(path))


def build_html(path: str) -> None:
    Path(path).write_text(
        f"<!doctype html><html><body><p>{MARKER_TEXT}</p></body></html>", encoding="utf-8"
    )


def build_markdown(path: str) -> None:
    Path(path).write_text(f"# Test\n\n{MARKER_TEXT}\n", encoding="utf-8")
