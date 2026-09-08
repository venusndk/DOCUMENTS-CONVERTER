"""
Tests for the Office/HTML/Markdown -> PDF capabilities
(documents_converter/converters/office_to_pdf.py, html_to_pdf.py,
markdown_to_pdf.py, _libreoffice.py -- Phase 9, master directive
numbering) and the reverse PDF -> Office direction
(pdf_to_docx.py, pdf_to_pptx.py -- Phase 10).

All gated by @requires_libreoffice: LibreOffice is not installed on
this project's own Windows dev machine (a deliberate choice -- see the
Phase 9 commit message -- not an oversight), so these skip locally and
run for real in Docker and CI, where it's apt-installed. Real
verification for both phases happened against an actual running Docker
container (see each phase's commit message), not only here.
"""

from __future__ import annotations

import sys

import fitz
import pytest

sys.path.insert(0, "tests/fixtures")

from conftest import requires_libreoffice
from documents_converter.api import config
from documents_converter.converters._libreoffice import (
    check_libreoffice_available,
    convert_via_libreoffice,
)
from documents_converter.converters.html_to_pdf import _convert as html_to_pdf_convert
from documents_converter.converters.markdown_to_pdf import _convert as markdown_to_pdf_convert
from documents_converter.converters.office_to_pdf import _convert as office_to_pdf_convert
from documents_converter.converters.pdf_to_docx import _convert as pdf_to_docx_convert
from documents_converter.converters.pdf_to_pptx import _convert as pdf_to_pptx_convert
from synthetic_office import MARKER_TEXT, build_docx, build_html, build_markdown, build_pptx, build_xlsx


def _build_text_pdf(path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    page.insert_text((36, 100), MARKER_TEXT)
    doc.save(str(path))
    doc.close()


def _pdf_text(path) -> str:
    doc = fitz.open(str(path))
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def test_check_libreoffice_available_reflects_reality():
    # Not gated by @requires_libreoffice on purpose -- this should give
    # a real, correct answer either way (True in Docker/CI, False here).
    from conftest import LIBREOFFICE_CMD

    assert check_libreoffice_available() == (LIBREOFFICE_CMD is not None)


def test_convert_via_libreoffice_raises_environment_error_when_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "documents_converter.converters._libreoffice.shutil.which", lambda _cmd: None
    )
    with pytest.raises(EnvironmentError):
        convert_via_libreoffice(tmp_path / "in.html", tmp_path / "out.pdf", soffice_cmd=None)


@requires_libreoffice
def test_docx_to_pdf_contains_the_source_text(tmp_path, libreoffice_cmd, monkeypatch):
    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    docx_path = tmp_path / "doc.docx"
    build_docx(docx_path)

    output_path = tmp_path / "out.pdf"
    office_to_pdf_convert(docx_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_xlsx_to_pdf_contains_the_source_text(tmp_path, libreoffice_cmd, monkeypatch):
    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    xlsx_path = tmp_path / "sheet.xlsx"
    build_xlsx(xlsx_path)

    output_path = tmp_path / "out.pdf"
    office_to_pdf_convert(xlsx_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_pptx_to_pdf_contains_the_source_text(tmp_path, libreoffice_cmd, monkeypatch):
    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    pptx_path = tmp_path / "slides.pptx"
    build_pptx(pptx_path)

    output_path = tmp_path / "out.pdf"
    office_to_pdf_convert(pptx_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_html_to_pdf_contains_the_source_text(tmp_path, libreoffice_cmd, monkeypatch):
    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    html_path = tmp_path / "page.html"
    build_html(html_path)

    output_path = tmp_path / "out.pdf"
    html_to_pdf_convert(html_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_markdown_to_pdf_contains_the_source_text(tmp_path, libreoffice_cmd, monkeypatch):
    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    md_path = tmp_path / "doc.md"
    build_markdown(md_path)

    output_path = tmp_path / "out.pdf"
    markdown_to_pdf_convert(md_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_concurrent_libreoffice_conversions_do_not_conflict(tmp_path, libreoffice_cmd, monkeypatch):
    """The specific known LibreOffice failure mode _libreoffice.py's own
    docstring documents -- multiple headless instances sharing a user
    profile fail against each other. This project's job queue runs
    conversions concurrently (api/app.py's _convert_executor), so this
    must actually work, not just look right in isolation."""
    import concurrent.futures

    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    html_path = tmp_path / "page.html"
    build_html(html_path)

    def _run(i):
        output_path = tmp_path / f"out_{i}.pdf"
        html_to_pdf_convert(html_path, output_path, progress=lambda _m: None)
        return output_path

    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as executor:
        results = list(executor.map(_run, range(4)))

    for output_path in results:
        assert output_path.exists()
        assert MARKER_TEXT in _pdf_text(output_path)


@requires_libreoffice
def test_pdf_to_docx_produces_a_valid_openable_docx(tmp_path, libreoffice_cmd, monkeypatch):
    """Reads the result back with python-docx (a different library than
    anything in the runtime conversion path -- LibreOffice wrote it,
    this just verifies it's a genuinely valid, openable file), not just
    checking the output path exists.

    Checks the document's raw XML for the marker text, not
    `document.paragraphs` -- confirmed directly against a real converted
    file (see pdf_to_docx.py's own docstring): LibreOffice's PDF import
    places recovered text in floating text-box shapes anchored at the
    PDF's original coordinates, not flowing body paragraphs, so
    python-docx's high-level `paragraphs` API (top-level body text only)
    genuinely does not see it even though the text is really there.
    """
    import zipfile

    from docx import Document

    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    pdf_path = tmp_path / "source.pdf"
    _build_text_pdf(pdf_path)

    output_path = tmp_path / "out.docx"
    pdf_to_docx_convert(pdf_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    Document(str(output_path))  # raises if this isn't a genuinely valid docx
    with zipfile.ZipFile(output_path) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
    assert MARKER_TEXT in document_xml


@requires_libreoffice
def test_pdf_to_pptx_produces_a_valid_openable_pptx_with_editable_text(
    tmp_path, libreoffice_cmd, monkeypatch
):
    """Unlike PDF -> DOCX (see that test's own note), LibreOffice's PDF ->
    Impress import puts recovered text into a real text-frame shape
    python-pptx's normal API reads directly -- confirmed against a real
    converted file, checked here rather than assumed to work the same
    way as the DOCX direction."""
    from pptx import Presentation

    monkeypatch.setattr(config, "LIBREOFFICE_CMD", libreoffice_cmd)
    pdf_path = tmp_path / "source.pdf"
    _build_text_pdf(pdf_path)

    output_path = tmp_path / "out.pptx"
    pdf_to_pptx_convert(pdf_path, output_path, progress=lambda _m: None)

    assert output_path.exists()
    prs = Presentation(str(output_path))
    assert len(prs.slides) >= 1
    all_text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert MARKER_TEXT in all_text
